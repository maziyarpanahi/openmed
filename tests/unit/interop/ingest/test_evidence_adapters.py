"""Synthetic structured-source evidence adapter tests."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pytest
from hypothesis import given
from hypothesis import strategies as st
from jsonschema.validators import validator_for
from openpyxl import Workbook

from openmed.interop.ingest import (
    EVIDENCE_ADAPTER_SCHEMA_NAMES,
    CDAEvidenceAdapter,
    DelimitedTableEvidenceAdapter,
    EvidenceAdapterContext,
    ExistingDocumentEvidenceAdapter,
    ExistingDocumentInput,
    FHIRR4EvidenceAdapter,
    HL7V2EvidenceAdapter,
    StructuredEvidenceAdapter,
    StructuredEvidenceQuarantine,
    StructuredEvidenceResult,
    TextEvidenceAdapter,
    XLSXEvidenceAdapter,
    load_all_evidence_adapter_schemas,
)
from openmed.multimodal.base import ExtractedDocument, SourceSpan
from openmed.structured.store import StoreState

T0 = "2026-01-02T03:04:05Z"
GOLDEN_FIXTURE = "tests/fixtures/interop/structured_evidence_golden.json"


@pytest.fixture
def context() -> EvidenceAdapterContext:
    return EvidenceAdapterContext(
        source_id="source_0123456789abcdef",
        subject_id="patient_0123456789abcdef",
        encounter_id="encounter_0123456789abcdef",
        recorded_at=T0,
    )


def _value(result) -> StructuredEvidenceResult:
    assert result.ok and isinstance(result.value, StructuredEvidenceResult)
    return result.value


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "SYNTHETIC Patient Name"
    first["A1"] = "label"
    first["B2"] = "café"
    second = workbook.create_sheet("Labs")
    second["C3"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_structured_evidence_golden_journey(
    context: EvidenceAdapterContext,
) -> None:
    fixture = json.loads(Path(GOLDEN_FIXTURE).read_text(encoding="utf-8"))
    adapters = {
        "text": TextEvidenceAdapter(),
        "fhir_r4": FHIRR4EvidenceAdapter(),
        "hl7v2": HL7V2EvidenceAdapter(),
        "cda": CDAEvidenceAdapter(),
        "csv": DelimitedTableEvidenceAdapter(),
    }

    assert fixture["schema_version"] == 1
    for case in fixture["cases"]:
        adapter = adapters[case["format"]]
        result = _value(adapter.adapt(case["source"], context))
        resolved = [
            adapter.resolve(case["source"], item).value for item in result.locators
        ]

        assert result.source_format == case["format"]
        assert {item.location_type for item in result.locators} == {
            case["expected_location_type"]
        }
        assert set(case["required_values"]).issubset(resolved)


def test_text_unicode_newlines_round_trip_and_replay_is_deterministic(
    context: EvidenceAdapterContext,
) -> None:
    source = "café\r\nمرحبا\n"
    adapter = TextEvidenceAdapter()

    first = _value(adapter.adapt(source, context))
    second = _value(adapter.adapt(source, context))

    assert first == second
    assert first.source_byte_size == len(source.encode("utf-8"))
    assert first.coordinate_convention == "unicode_text_v1"
    assert len(first.locators) == 1
    resolved = adapter.resolve(source, first.locators[0])
    assert resolved.ok and resolved.value == source
    assert source not in first.to_json()
    assert source not in repr(first)


@given(
    st.text(
        alphabet=st.characters(blacklist_categories=("Cs",)),
        min_size=1,
        max_size=128,
    )
)
def test_text_property_round_trip_preserves_code_points(source: str) -> None:
    context = EvidenceAdapterContext(
        source_id="source_0123456789abcdef",
        recorded_at=T0,
    )
    adapter = TextEvidenceAdapter()

    result = _value(adapter.adapt(source, context))

    assert len(result.locators) == 1
    assert adapter.resolve(source, result.locators[0]).value == source


def test_source_limits_and_invalid_utf8_return_value_free_typed_outcomes(
    context: EvidenceAdapterContext,
) -> None:
    denied = TextEvidenceAdapter(max_source_bytes=1).adapt(b"ab", context)
    invalid_source = b"synthetic-canary-\xff"
    failed = TextEvidenceAdapter().adapt(invalid_source, context)

    assert denied.state is StoreState.DENIED
    assert denied.code == "source_limit_exceeded"
    assert isinstance(denied.value, StructuredEvidenceQuarantine)
    assert failed.state is StoreState.FAILURE
    assert isinstance(failed.value, StructuredEvidenceQuarantine)
    assert failed.value.source_byte_size == len(invalid_source)
    assert "synthetic-canary" not in failed.value.to_json()


def test_locator_refuses_different_source_bytes(
    context: EvidenceAdapterContext,
) -> None:
    adapter = TextEvidenceAdapter()
    result = _value(adapter.adapt("first source", context))

    mismatch = adapter.resolve("second source", result.locators[0])

    assert mismatch.state is StoreState.CONFLICT
    assert mismatch.code == "source_digest_mismatch"


def test_existing_document_preserves_normalized_and_page_coordinates(
    context: EvidenceAdapterContext,
) -> None:
    document = ExtractedDocument(
        text="alpha βeta",
        spans=(
            SourceSpan(start=0, end=5, page=0),
            SourceSpan(start=6, end=10, page=1, bbox=(10.0, 20.0, 30.0, 40.0)),
        ),
    )
    source = ExistingDocumentInput(
        document=document,
        source_bytes=b"synthetic document bytes",
        bbox_coordinate_space="points",
    )
    adapter = ExistingDocumentEvidenceAdapter()

    result = _value(adapter.adapt(source, context))

    assert {item.location_type for item in result.locators} == {
        "page_box",
        "text_span",
    }
    assert sorted(adapter.resolve(source, item).value for item in result.locators) == [
        "alpha",
        "βeta",
    ]


def test_existing_document_rejects_ambiguous_box_coordinates(
    context: EvidenceAdapterContext,
) -> None:
    document = ExtractedDocument(
        text="synthetic",
        spans=(SourceSpan(start=0, end=9, bbox=(0.0, 0.0, 1.0, 1.0)),),
    )

    result = ExistingDocumentEvidenceAdapter().adapt(
        ExistingDocumentInput(document=document, source_bytes=b"source"),
        context,
    )

    assert result.state is StoreState.CONFLICT
    assert result.code == "coordinate_space_ambiguous"
    assert isinstance(result.value, StructuredEvidenceQuarantine)


def test_fhir_r4_every_scalar_round_trips_to_its_json_pointer(
    context: EvidenceAdapterContext,
) -> None:
    resource = {
        "resourceType": "Observation",
        "status": "final",
        "code": {"text": "synthetic café"},
        "component": [{"valueString": "مرحبا"}],
    }
    source = json.dumps(resource, ensure_ascii=False, separators=(",", ":"))
    adapter = FHIRR4EvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    resolved = {
        item.location["pointer"]: adapter.resolve(source, item).value
        for item in result.locators
    }

    assert resolved == {
        "/code/text": "synthetic café",
        "/component/0/valueString": "مرحبا",
        "/resourceType": "Observation",
        "/status": "final",
    }
    assert "synthetic café" not in result.to_json()


def test_fhir_unsupported_version_and_field_limit_are_typed(
    context: EvidenceAdapterContext,
) -> None:
    source = '{"resourceType":"Observation","status":"final"}'

    unsupported = FHIRR4EvidenceAdapter(source_version="5.0.0").adapt(source, context)
    partial = FHIRR4EvidenceAdapter(max_fields=1).adapt(source, context)

    assert unsupported.state is StoreState.UNSUPPORTED
    assert unsupported.code == "fhir_version_unsupported"
    assert isinstance(unsupported.value, StructuredEvidenceQuarantine)
    assert partial.state is StoreState.PARTIAL
    assert partial.code == "field_limit_exceeded"


def test_fhir_duplicate_keys_and_nonfinite_values_quarantine(
    context: EvidenceAdapterContext,
) -> None:
    adapter = FHIRR4EvidenceAdapter()

    duplicate = adapter.adapt(
        '{"resourceType":"Observation","status":"a","status":"b"}',
        context,
    )
    nonfinite = adapter.adapt(
        '{"resourceType":"Observation","valueDecimal":NaN}',
        context,
    )

    assert duplicate.state is StoreState.FAILURE
    assert duplicate.code == "fhir_json_malformed"
    assert nonfinite.state is StoreState.FAILURE


def test_hl7_custom_delimiters_and_components_round_trip(
    context: EvidenceAdapterContext,
) -> None:
    source = (
        "MSH*$%!?*SYNTHETIC*LAB*OPENMED*LOCAL*20260102030405**"
        "ORU$R01*MSG0001*P*2.5\r"
        "PID*1**opaque$local**synthetic\r"
    )
    adapter = HL7V2EvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    by_path = {item.location["path"]: item for item in result.locators}

    assert adapter.resolve(source, by_path["PID.3.1[1]"]).value == "opaque"
    assert adapter.resolve(source, by_path["PID.3.2[1]"]).value == "local"
    assert adapter.resolve(source, by_path["MSH.12.1[1]"]).value == "2.5"


def test_hl7_unsupported_version_and_repetition_are_quarantined(
    context: EvidenceAdapterContext,
) -> None:
    unsupported_source = (
        "MSH|^~\\&|SYNTHETIC|LAB|OPENMED|LOCAL|20260102030405||ORU^R01|MSG0001|P|2.9\r"
    )
    repeated_source = (
        "MSH|^~\\&|SYNTHETIC|LAB|OPENMED|LOCAL|20260102030405||"
        "ORU^R01|MSG0001|P|2.5\rPID|1||first~second\r"
    )
    adapter = HL7V2EvidenceAdapter()

    unsupported = adapter.adapt(unsupported_source, context)
    repeated = adapter.adapt(repeated_source, context)

    assert unsupported.state is StoreState.UNSUPPORTED
    assert unsupported.code == "hl7_version_unsupported"
    assert repeated.state is StoreState.CONFLICT
    assert repeated.code == "hl7_coordinate_ambiguous"


def test_cda_sections_text_and_attributes_round_trip(
    context: EvidenceAdapterContext,
) -> None:
    source = b"""<?xml version="1.0" encoding="UTF-8"?>
<ClinicalDocument xmlns="urn:hl7-org:v3">
  <component><structuredBody><component><section>
    <code code="synthetic-code"/>
    <title>Results</title>
    <text><paragraph>cafe synthetic</paragraph></text>
  </section></component></structuredBody></component>
</ClinicalDocument>"""
    adapter = CDAEvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    resolved = [adapter.resolve(source, item).value for item in result.locators]

    assert "synthetic-code" in resolved
    assert "Results" in resolved
    assert "cafe synthetic" in resolved
    section_locators = [item for item in result.locators if "section" in item.location]
    assert section_locators
    assert {item.location["section"] for item in section_locators} == {1}
    assert "synthetic-code" not in result.to_json()


def test_cda_unsafe_declaration_and_version_are_typed_quarantine(
    context: EvidenceAdapterContext,
) -> None:
    source = b'<ClinicalDocument xmlns="urn:hl7-org:v3"/>'
    unsafe = b'<!DOCTYPE ClinicalDocument><ClinicalDocument xmlns="urn:hl7-org:v3"/>'

    unsupported = CDAEvidenceAdapter(source_version="3.0").adapt(source, context)
    denied = CDAEvidenceAdapter().adapt(unsafe, context)

    assert unsupported.state is StoreState.UNSUPPORTED
    assert unsupported.code == "cda_version_unsupported"
    assert denied.state is StoreState.DENIED
    assert denied.code == "cda_declaration_unsafe"


def test_csv_bom_unicode_quoted_newline_and_indexes_round_trip(
    context: EvidenceAdapterContext,
) -> None:
    source = b'\xef\xbb\xbfname,note\r\nsynthetic,"caf\xc3\xa9\nline"\r\n'
    adapter = DelimitedTableEvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    values = {
        (item.location["row"], item.location["column"]): adapter.resolve(
            source, item
        ).value
        for item in result.locators
    }

    assert values == {
        (1, 1): "name",
        (1, 2): "note",
        (2, 1): "synthetic",
        (2, 2): "café\nline",
    }


@given(
    st.lists(
        st.lists(
            st.text(
                alphabet=st.characters(blacklist_categories=("Cs",)),
                max_size=24,
            ),
            min_size=1,
            max_size=5,
        ),
        min_size=1,
        max_size=5,
    )
)
def test_csv_property_round_trip_preserves_logical_cells(rows: list[list[str]]) -> None:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerows(rows)
    source = output.getvalue()
    context = EvidenceAdapterContext(
        source_id="source_0123456789abcdef",
        recorded_at=T0,
    )
    adapter = DelimitedTableEvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    resolved = [adapter.resolve(source, item).value for item in result.locators]

    assert sorted(resolved) == sorted(value for row in rows for value in row)


def test_csv_custom_delimiter_and_malformed_input(
    context: EvidenceAdapterContext,
) -> None:
    tab = DelimitedTableEvidenceAdapter(delimiter="\t")
    result = _value(tab.adapt("a\tb\r\n1\t2\r\n", context))
    assert len(result.locators) == 4

    malformed = DelimitedTableEvidenceAdapter().adapt('a,"unterminated', context)
    assert malformed.state is StoreState.FAILURE
    assert malformed.code == "csv_malformed"


def test_xlsx_uses_opaque_sheet_coordinates_and_round_trips_cells(
    context: EvidenceAdapterContext,
) -> None:
    source = _xlsx_bytes()
    adapter = XLSXEvidenceAdapter()

    result = _value(adapter.adapt(source, context))
    locations = {
        (item.location["sheet"], item.location["row"], item.location["column"]): item
        for item in result.locators
    }

    assert set(locations) == {
        ("sheet_0001", 1, 1),
        ("sheet_0001", 2, 2),
        ("sheet_0002", 3, 3),
    }
    assert adapter.resolve(source, locations[("sheet_0001", 2, 2)]).value == "café"
    assert adapter.resolve(source, locations[("sheet_0002", 3, 3)]).value == "=1+1"
    assert "SYNTHETIC Patient Name" not in result.to_json()
    assert "SYNTHETIC Patient Name" not in repr(result)


def test_public_results_round_trip_and_validate_bundled_schemas(
    context: EvidenceAdapterContext,
) -> None:
    result = _value(TextEvidenceAdapter().adapt("synthetic", context))
    quarantine_result = FHIRR4EvidenceAdapter(source_version="5.0.0").adapt(
        '{"resourceType":"Observation"}', context
    )
    assert isinstance(quarantine_result.value, StructuredEvidenceQuarantine)
    quarantine = quarantine_result.value
    records = {"result": result, "quarantine": quarantine}
    schemas = load_all_evidence_adapter_schemas()

    assert set(schemas) == set(EVIDENCE_ADAPTER_SCHEMA_NAMES)
    assert StructuredEvidenceResult.from_json(result.to_json()) == result
    assert StructuredEvidenceQuarantine.from_json(quarantine.to_json()) == quarantine
    for name, record in records.items():
        schema = schemas[name]
        validator = validator_for(schema)
        validator.check_schema(schema)
        assert schema["schema_version"] == 1
        assert not tuple(validator(schema).iter_errors(record.to_dict()))


def test_adapters_satisfy_public_protocol() -> None:
    assert isinstance(TextEvidenceAdapter(), StructuredEvidenceAdapter)
    assert isinstance(FHIRR4EvidenceAdapter(), StructuredEvidenceAdapter)
    assert isinstance(HL7V2EvidenceAdapter(), StructuredEvidenceAdapter)
    assert isinstance(CDAEvidenceAdapter(), StructuredEvidenceAdapter)
    assert isinstance(DelimitedTableEvidenceAdapter(), StructuredEvidenceAdapter)
    assert isinstance(XLSXEvidenceAdapter(), StructuredEvidenceAdapter)
