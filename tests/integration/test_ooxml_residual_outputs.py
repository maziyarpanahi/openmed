"""Offline synthetic Office write-boundary tests; no model or network access."""

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

from openmed.multimodal.documents_docx import extract_docx, write_redacted_docx
from openmed.multimodal.ooxml_residual import OoxmlResidualError
from openmed.multimodal.pptx import extract_pptx, write_redacted_pptx
from openmed.multimodal.xlsx import redact_xlsx

pytestmark = pytest.mark.integration
SENTINEL = "SYNTHETIC_PRIVATE_患者_3810"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _fixture(path: Path, kind: str):
    if kind == "docx":
        docx = pytest.importorskip("docx")
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Patient ")
        paragraph.add_run("Jane")
        paragraph.add_run(" Doe")
        document.save(path)
    elif kind == "pptx":
        pptx = pytest.importorskip("pptx")
        from pptx.util import Inches

        deck = pptx.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        ).text = "Patient Jane Doe"
        deck.save(path)
    else:
        openpyxl = pytest.importorskip("openpyxl")
        workbook = openpyxl.Workbook()
        workbook.active.append(["patient_name", "count"])
        workbook.active.append(["Jane Doe", 2])
        workbook.save(path)
        workbook.close()


def _rewrite(path, mutate):
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    mutate(parts)
    with zipfile.ZipFile(path, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)


def _write(source, output, kind):
    if kind == "xlsx":
        return redact_xlsx(source, output, cell_deidentifier=lambda value: value)
    extract, write = (
        (extract_docx, write_redacted_docx)
        if kind == "docx"
        else (extract_pptx, write_redacted_pptx)
    )
    document = extract(source)
    start = document.text.index("Jane Doe")
    return write(
        source, output, [{"start": start, "end": start + 8, "label": "PERSON"}]
    )


def _text(parts, member, tag, value=SENTINEL):
    root = ET.fromstring(parts[member])
    ET.SubElement(root, tag).text = value
    parts[member] = ET.tostring(root)


def _hidden(parts, kind, category):
    main = {
        "docx": "word/document.xml",
        "xlsx": "xl/worksheets/sheet1.xml",
        "pptx": "ppt/slides/slide1.xml",
    }[kind]
    prefix = {"docx": "word", "xlsx": "xl", "pptx": "ppt"}[kind]
    if category in {"core_properties", "app_properties", "custom_properties"}:
        name = {
            "core_properties": "core",
            "app_properties": "app",
            "custom_properties": "custom",
        }[category]
        parts[f"docProps/{name}.xml"] = (
            f'<properties author="{SENTINEL}">{SENTINEL}</properties>'.encode()
        )
    elif category == "custom_xml":
        parts["customXml/private.xml"] = f"<data>{SENTINEL}</data>".encode()
    elif category == "alt_text":
        root = ET.fromstring(parts[main])
        ET.SubElement(root, f"{{{A}}}cNvPr", {"descr": SENTINEL, "title": SENTINEL})
        parts[main] = ET.tostring(root)
    elif category == "tracked_changes":
        root = ET.fromstring(parts[main])
        deletion = ET.SubElement(
            root.find(f"{{{W}}}body"), f"{{{W}}}del", {f"{{{W}}}author": SENTINEL}
        )
        ET.SubElement(deletion, f"{{{W}}}delText").text = SENTINEL
        parts[main] = ET.tostring(root)
    elif category == "defined_names":
        _text(parts, "xl/workbook.xml", f"{{{S}}}definedName")
    elif category == "hidden_sheets":
        root = ET.fromstring(parts["xl/workbook.xml"])
        root.find(f".//{{{S}}}sheet").set("state", "veryHidden")
        parts["xl/workbook.xml"] = ET.tostring(root)
    elif category == "headers_footers":
        if kind == "docx":
            parts["word/header99.xml"] = f"<header>{SENTINEL}</header>".encode()
        elif kind == "xlsx":
            root = ET.fromstring(parts[main])
            header = ET.SubElement(root, f"{{{S}}}headerFooter")
            ET.SubElement(header, f"{{{S}}}oddHeader").text = SENTINEL
            parts[main] = ET.tostring(root)
        else:
            _text(parts, main, "footer")
    else:
        paths = {
            "comments_notes": f"{prefix}/comments99.xml",
            "pivot_caches": f"{prefix}/pivotCache/private.xml",
            "external_links": f"{prefix}/externalLinks/private.xml",
            "embedded_objects": f"{prefix}/embeddings/private.bin",
            "uncovered_part": f"{prefix}/private.xml",
        }
        parts[paths[category]] = (
            f'<private author="{SENTINEL}">{SENTINEL}</private>'.encode()
        )


CASES = [
    (kind, category)
    for kind in ("docx", "xlsx", "pptx")
    for category in (
        "core_properties",
        "app_properties",
        "custom_properties",
        "custom_xml",
        "alt_text",
        "comments_notes",
        "headers_footers",
        "pivot_caches",
        "external_links",
        "embedded_objects",
        "uncovered_part",
    )
] + [("docx", "tracked_changes"), ("xlsx", "defined_names"), ("xlsx", "hidden_sheets")]


@pytest.mark.parametrize("kind,category", CASES)
def test_hidden_sentinel_removed_or_typed_refusal(tmp_path, kind, category):
    source = tmp_path / f"source.{kind}"
    output = tmp_path / f"output.{kind}"
    _fixture(source, kind)
    _rewrite(source, lambda parts: _hidden(parts, kind, category))
    original = source.read_bytes()
    output.write_bytes(b"existing destination")
    try:
        _write(source, output, kind)
    except OoxmlResidualError as error:
        assert output.read_bytes() == b"existing destination"
        assert SENTINEL not in str(error) + repr(error.findings)
        assert str(source) not in str(error) + repr(error.findings)
        assert (
            category in {item.category for item in error.findings}
            or category == "headers_footers"
        )
    else:
        with zipfile.ZipFile(output) as archive:
            assert all(
                SENTINEL.encode() not in archive.read(name)
                for name in archive.namelist()
            )
        assert SENTINEL.encode() not in output.read_bytes()
        if kind != "xlsx":
            extract = extract_docx if kind == "docx" else extract_pptx
            assert "Patient [PERSON]" in extract(output).text
    assert source.read_bytes() == original
    assert not list(tmp_path.glob(".ooxml-*"))


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_stream_and_in_place_refusal_preserve_destinations(tmp_path, kind):
    source = tmp_path / f"source.{kind}"
    _fixture(source, kind)
    _rewrite(source, lambda parts: _hidden(parts, kind, "uncovered_part"))
    original = source.read_bytes()
    stream = io.BytesIO(b"existing stream")
    with pytest.raises(OoxmlResidualError):
        _write(source, stream, kind)
    assert stream.getvalue() == b"existing stream"
    with pytest.raises(OoxmlResidualError):
        _write(source, source, kind)
    assert source.read_bytes() == original


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_verifier_gates_staged_output(tmp_path, monkeypatch, kind):
    import openmed.multimodal.ooxml_residual as residual

    source = tmp_path / f"source.{kind}"
    output = tmp_path / f"output.{kind}"
    _fixture(source, kind)
    output.write_bytes(b"existing destination")
    sanitize = residual._sanitize
    calls = 0

    def inject_after_save(data):
        nonlocal calls
        calls += 1
        safe = sanitize(data)
        if calls == 2:
            buffer = io.BytesIO(safe)
            with zipfile.ZipFile(buffer, "a") as archive:
                archive.writestr("uncovered.xml", f"<private>{SENTINEL}</private>")
            return buffer.getvalue()
        return safe

    monkeypatch.setattr(residual, "_sanitize", inject_after_save)
    with pytest.raises(OoxmlResidualError, match="uncovered_part"):
        _write(source, output, kind)
    assert output.read_bytes() == b"existing destination"
    assert calls == 2


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_nested_text_in_existing_run_is_not_declared_covered(tmp_path, kind):
    source = tmp_path / f"source.{kind}"
    output = tmp_path / f"output.{kind}"
    _fixture(source, kind)

    def nest(parts):
        member = "word/document.xml" if kind == "docx" else "ppt/slides/slide1.xml"
        ns = W if kind == "docx" else A
        root = ET.fromstring(parts[member])
        run = root.find(f".//{{{ns}}}r")
        unknown = ET.SubElement(run, f"{{{ns}}}uncovered")
        ET.SubElement(unknown, f"{{{ns}}}t").text = SENTINEL
        parts[member] = ET.tostring(root)

    _rewrite(source, nest)
    with pytest.raises(OoxmlResidualError, match="uncovered_text"):
        _write(source, output, kind)
    assert not output.exists()
