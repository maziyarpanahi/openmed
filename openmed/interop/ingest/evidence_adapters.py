"""Offline adapters from structured sources to Journey evidence coordinates."""

from __future__ import annotations

import csv
import importlib.util
import io
import json
import math
import re
from collections.abc import Mapping
from dataclasses import dataclass
from typing import Any, Protocol, runtime_checkable
from xml.etree import ElementTree as ET

from openmed.clinical.journey_contracts import (
    ClinicalArtifact,
    EvidenceLocator,
    derived_opaque_id,
    sha256_digest,
)
from openmed.interop.hl7v2 import HL7Message
from openmed.multimodal.base import ExtractedDocument
from openmed.structured.store import StoreResult, StoreState

from .evidence_contracts import (
    StructuredEvidenceQuarantine,
    StructuredEvidenceResult,
)

DEFAULT_MAX_SOURCE_BYTES = 64 * 1024 * 1024
DEFAULT_MAX_FIELDS = 100_000

_OPAQUE_ID_RE = re.compile(r"^[a-z][a-z0-9_]{0,31}_[A-Za-z0-9_-]{16,128}$")
_TIMESTAMP_RE = re.compile(
    r"^[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}"
    r"(?:\.[0-9]+)?(?:Z|[+-][0-9]{2}:[0-9]{2})$"
)
_SAFE_VERSION_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,63}$")
_FHIR_KEY_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,127}$")
_HL7_PATH_RE = re.compile(
    r"^(?P<segment>[A-Z0-9][A-Z0-9_-]*)\."
    r"(?P<field>[1-9][0-9]*)\.(?P<component>[1-9][0-9]*)"
    r"\[(?P<occurrence>[1-9][0-9]*)\]$"
)
_XML_SEGMENT_RE = re.compile(
    r"^(?P<name>[A-Za-z_][A-Za-z0-9_.-]*)\[(?P<index>[1-9][0-9]*)\]$"
)
_UNSAFE_XML_RE = re.compile(rb"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)
_CDA_NAMESPACE = "urn:hl7-org:v3"
_XSI_NAMESPACE = "http://www.w3.org/2001/XMLSchema-instance"

_SUPPORTED_FHIR_VERSIONS = frozenset({"4.0", "4.0.1"})
_SUPPORTED_HL7_VERSIONS = frozenset(
    {
        "2.3",
        "2.3.1",
        "2.4",
        "2.5",
        "2.5.1",
        "2.6",
        "2.7",
        "2.7.1",
        "2.8",
        "2.8.1",
        "2.8.2",
    }
)
_SUPPORTED_CDA_VERSIONS = frozenset({"2", "2.0"})


@dataclass(frozen=True, slots=True)
class EvidenceAdapterContext:
    """Opaque identity and time metadata shared by every adapter."""

    source_id: str
    recorded_at: str
    subject_id: str | None = None
    encounter_id: str | None = None

    def __post_init__(self) -> None:
        for field_name in ("source_id", "subject_id", "encounter_id"):
            value = getattr(self, field_name)
            if value is not None and _OPAQUE_ID_RE.fullmatch(value) is None:
                raise ValueError(f"{field_name} must be an opaque identifier")
        if _TIMESTAMP_RE.fullmatch(self.recorded_at) is None:
            raise ValueError("recorded_at must be a timezone-aware ISO timestamp")


@dataclass(frozen=True, slots=True)
class ExistingDocumentInput:
    """Existing normalized document plus immutable original source bytes."""

    document: ExtractedDocument
    source_bytes: bytes
    bbox_coordinate_space: str | None = None

    def __post_init__(self) -> None:
        if not isinstance(self.document, ExtractedDocument):
            raise TypeError("document must be an ExtractedDocument")
        if not isinstance(self.source_bytes, bytes):
            raise TypeError("source_bytes must be bytes")
        if self.bbox_coordinate_space not in {None, "normalized", "pixels", "points"}:
            raise ValueError("bbox_coordinate_space is unsupported")


EvidenceAdapterOutcome = StructuredEvidenceResult | StructuredEvidenceQuarantine
LocatorSpec = tuple[str, Mapping[str, Any], Mapping[str, Any]]


@runtime_checkable
class StructuredEvidenceAdapter(Protocol):
    """Backend-neutral structured-source evidence adapter."""

    parser_id: str
    parser_version: str

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Return exact evidence metadata or a typed quarantine."""

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one locator against caller-held source content."""


class TextEvidenceAdapter:
    """Adapt UTF-8 text without newline or Unicode normalization."""

    parser_id = "openmed.evidence.text"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
    ) -> None:
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Create one full-text locator while preserving exact code-point offsets."""

        coerced = _coerce_utf8(source)
        if coerced is None:
            return _quarantine(
                _coerce_bytes(source) or b"",
                context,
                source_format="text",
                source_version="plain",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="text_source_invalid",
            )
        payload, text = coerced
        limited = _limit_quarantine(
            payload,
            context,
            source_format="text",
            source_version="plain",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        specs: tuple[LocatorSpec, ...] = ()
        if text:
            specs = (("text_span", {"start": 0, "end": len(text)}, {}),)
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="text",
            source_version="plain",
            media_type="text/plain",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="unicode_text_v1",
            normalization_transform={
                "operation": "identity",
                "version": "1.0.0",
                "encoding": "utf8",
                "newline": "preserve",
                "unicode": "preserve",
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Return the exact source substring selected by a text locator."""

        coerced = _coerce_utf8(source)
        if coerced is None or locator.location_type != "text_span":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, coerced[0]):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        text = coerced[1]
        start = int(locator.location["start"])
        end = int(locator.location["end"])
        if end > len(text):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_out_of_bounds")
        return StoreResult.success(text[start:end])


class ExistingDocumentEvidenceAdapter:
    """Adapt an existing normalized document and its source-coordinate spans."""

    parser_id = "openmed.evidence.document"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Preserve normalized spans and explicit original page geometry."""

        if not isinstance(source, ExistingDocumentInput):
            return _quarantine(
                b"",
                context,
                source_format="document",
                source_version="provided",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="document_source_invalid",
            )
        limited = _limit_quarantine(
            source.source_bytes,
            context,
            source_format="document",
            source_version="provided",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        if len(source.document.spans) > self.max_fields:
            return _field_limit_quarantine(
                source.source_bytes,
                context,
                source_format="document",
                source_version="provided",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=len(source.document.spans),
            )
        specs: list[tuple[str, Mapping[str, Any], Mapping[str, Any]]] = []
        previous_end = 0
        for span in source.document.spans:
            if (
                type(span.start) is not int
                or type(span.end) is not int
                or span.start < previous_end
                or span.end <= span.start
                or span.end > len(source.document.text)
            ):
                return _quarantine(
                    source.source_bytes,
                    context,
                    source_format="document",
                    source_version="provided",
                    parser_id=self.parser_id,
                    parser_version=self.parser_version,
                    state=StoreState.CONFLICT,
                    classification="ambiguous",
                    reason_code="document_span_ambiguous",
                    candidate_count=len(specs),
                )
            previous_end = span.end
            transform: dict[str, Any] = {
                "operation": "provided_document_projection",
                "version": "1.0.0",
                "normalized_start": span.start,
                "normalized_end": span.end,
                "source_page": span.page + 1,
            }
            if span.bbox is None:
                specs.append(
                    (
                        "text_span",
                        {"start": span.start, "end": span.end},
                        transform,
                    )
                )
                continue
            if source.bbox_coordinate_space is None:
                return _quarantine(
                    source.source_bytes,
                    context,
                    source_format="document",
                    source_version="provided",
                    parser_id=self.parser_id,
                    parser_version=self.parser_version,
                    state=StoreState.CONFLICT,
                    classification="ambiguous",
                    reason_code="coordinate_space_ambiguous",
                    candidate_count=len(specs),
                )
            transform["source_coordinate_space"] = source.bbox_coordinate_space
            specs.append(
                (
                    "page_box",
                    {
                        "page": span.page + 1,
                        "box": list(span.bbox),
                        "coordinate_space": source.bbox_coordinate_space,
                    },
                    transform,
                )
            )
        normalized = source.document.text.encode("utf-8")
        return _success(
            source_bytes=source.source_bytes,
            artifact_bytes=source.source_bytes,
            context=context,
            source_format="document",
            source_version="provided",
            media_type="application/octet-stream",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="existing_document_v1",
            normalization_transform={
                "operation": "provided_document_projection",
                "version": "1.0.0",
                "encoding": "utf8",
                "newline": "provided",
                "unicode": "provided",
                "normalized_digest": sha256_digest(normalized),
            },
            locator_specs=tuple(specs),
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Return normalized text selected by a document locator."""

        if not isinstance(source, ExistingDocumentInput):
            return StoreResult.outcome(StoreState.FAILURE, "document_source_invalid")
        if not _locator_matches_source(locator, source.source_bytes):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        transform = locator.transform
        start = transform.get("normalized_start")
        end = transform.get("normalized_end")
        if type(start) is not int or type(end) is not int:
            if locator.location_type != "text_span":
                return StoreResult.outcome(
                    StoreState.UNSUPPORTED, "locator_unsupported"
                )
            start = locator.location["start"]
            end = locator.location["end"]
        if start < 0 or end <= start or end > len(source.document.text):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_out_of_bounds")
        return StoreResult.success(source.document.text[start:end])


class FHIRR4EvidenceAdapter:
    """Adapt FHIR R4 JSON scalar fields to RFC 6901 pointers."""

    parser_id = "openmed.evidence.fhir_r4"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        source_version: str = "4.0.1",
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        self.source_version = source_version
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Parse one FHIR R4 resource and emit a locator for every scalar."""

        coerced = _coerce_utf8(source)
        if coerced is None:
            return _quarantine(
                _coerce_bytes(source) or b"",
                context,
                source_format="fhir_r4",
                source_version=_safe_version(self.source_version),
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="fhir_source_invalid",
            )
        payload, text = coerced
        limited = _limit_quarantine(
            payload,
            context,
            source_format="fhir_r4",
            source_version=_safe_version(self.source_version),
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        if self.source_version not in _SUPPORTED_FHIR_VERSIONS:
            return _quarantine(
                payload,
                context,
                source_format="fhir_r4",
                source_version=_safe_version(self.source_version),
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.UNSUPPORTED,
                classification="unsupported",
                reason_code="fhir_version_unsupported",
            )
        try:
            resource = _strict_json(text)
            if not isinstance(resource, Mapping) or not isinstance(
                resource.get("resourceType"), str
            ):
                raise ValueError("invalid resource")
            pointers = _json_leaf_pointers(resource)
        except (TypeError, ValueError, json.JSONDecodeError):
            return _quarantine(
                payload,
                context,
                source_format="fhir_r4",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="fhir_json_malformed",
            )
        if len(pointers) > self.max_fields:
            return _field_limit_quarantine(
                payload,
                context,
                source_format="fhir_r4",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=len(pointers),
            )
        specs: tuple[LocatorSpec, ...] = tuple(
            ("json_pointer", {"pointer": pointer}, {}) for pointer in pointers
        )
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="fhir_r4",
            source_version=self.source_version,
            media_type="application/fhir+json",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="fhir_json_pointer_v1",
            normalization_transform={
                "operation": "parse_only",
                "version": "1.0.0",
                "encoding": "utf8",
                "newline": "preserve",
                "unicode": "preserve",
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one FHIR JSON pointer without changing its value."""

        coerced = _coerce_utf8(source)
        if coerced is None or locator.location_type != "json_pointer":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, coerced[0]):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        try:
            resource = _strict_json(coerced[1])
            return StoreResult.success(
                _resolve_json_pointer(resource, str(locator.location["pointer"]))
            )
        except (KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_unresolvable")


class HL7V2EvidenceAdapter:
    """Adapt HL7 v2 fields and components using message-native delimiters."""

    parser_id = "openmed.evidence.hl7v2"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Parse a message and preserve one-based field/component coordinates."""

        coerced = _coerce_utf8(source)
        if coerced is None:
            return _quarantine(
                _coerce_bytes(source) or b"",
                context,
                source_format="hl7v2",
                source_version="unknown",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="hl7_source_invalid",
            )
        payload, text = coerced
        limited = _limit_quarantine(
            payload,
            context,
            source_format="hl7v2",
            source_version="unknown",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        try:
            message = HL7Message.parse(text)
            version = message.segments[0].get_field(12) or "unknown"
        except (TypeError, ValueError):
            return _quarantine(
                payload,
                context,
                source_format="hl7v2",
                source_version="unknown",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="hl7_message_malformed",
            )
        if version not in _SUPPORTED_HL7_VERSIONS:
            return _quarantine(
                payload,
                context,
                source_format="hl7v2",
                source_version=_safe_version(version),
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.UNSUPPORTED,
                classification="unsupported",
                reason_code="hl7_version_unsupported",
            )
        try:
            paths = _hl7_component_paths(message)
        except ValueError:
            return _quarantine(
                payload,
                context,
                source_format="hl7v2",
                source_version=version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.CONFLICT,
                classification="ambiguous",
                reason_code="hl7_coordinate_ambiguous",
            )
        if len(paths) > self.max_fields:
            return _field_limit_quarantine(
                payload,
                context,
                source_format="hl7v2",
                source_version=version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=len(paths),
            )
        specs: tuple[LocatorSpec, ...] = tuple(
            ("message_field", {"path": path}, {}) for path in paths
        )
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="hl7v2",
            source_version=version,
            media_type="application/hl7-v2",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="hl7v2_field_component_v1",
            normalization_transform={
                "operation": "parse_only",
                "version": "1.0.0",
                "encoding": "utf8",
                "newline": "preserve",
                "delimiter_mode": "message_native",
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one field/component path with the message delimiter set."""

        coerced = _coerce_utf8(source)
        if coerced is None or locator.location_type != "message_field":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, coerced[0]):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        try:
            message = HL7Message.parse(coerced[1])
            value = _resolve_hl7_component(message, str(locator.location["path"]))
            return StoreResult.success(value)
        except (IndexError, TypeError, ValueError):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_unresolvable")


class CDAEvidenceAdapter:
    """Adapt CDA R2 XML text and attributes to indexed document paths."""

    parser_id = "openmed.evidence.cda"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        source_version: str = "2.0",
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        self.source_version = source_version
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Parse safe CDA XML and preserve section-aware element paths."""

        payload = _coerce_bytes(source)
        if payload is None:
            payload = b""
            return _quarantine(
                payload,
                context,
                source_format="cda",
                source_version=_safe_version(self.source_version),
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="cda_source_invalid",
            )
        limited = _limit_quarantine(
            payload,
            context,
            source_format="cda",
            source_version=_safe_version(self.source_version),
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        if self.source_version not in _SUPPORTED_CDA_VERSIONS:
            return _quarantine(
                payload,
                context,
                source_format="cda",
                source_version=_safe_version(self.source_version),
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.UNSUPPORTED,
                classification="unsupported",
                reason_code="cda_version_unsupported",
            )
        if _UNSAFE_XML_RE.search(payload):
            return _quarantine(
                payload,
                context,
                source_format="cda",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.DENIED,
                classification="unsafe",
                reason_code="cda_declaration_unsafe",
            )
        try:
            root = ET.fromstring(payload)
            if (
                _local_name(root.tag) != "ClinicalDocument"
                or _namespace(root.tag) != _CDA_NAMESPACE
            ):
                raise ValueError("invalid root")
            paths = _cda_value_paths(root)
        except (ET.ParseError, TypeError, ValueError):
            return _quarantine(
                payload,
                context,
                source_format="cda",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="cda_xml_malformed",
            )
        if any(len(path) > 512 for path, _section in paths):
            return _quarantine(
                payload,
                context,
                source_format="cda",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.CONFLICT,
                classification="ambiguous",
                reason_code="cda_path_unsupported",
                candidate_count=len(paths),
            )
        if len(paths) > self.max_fields:
            return _field_limit_quarantine(
                payload,
                context,
                source_format="cda",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=len(paths),
            )
        specs: tuple[LocatorSpec, ...] = tuple(
            (
                "document_path",
                {"path": path} | ({"section": section} if section is not None else {}),
                {},
            )
            for path, section in paths
        )
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="cda",
            source_version=self.source_version,
            media_type="application/xml",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="cda_document_path_v1",
            normalization_transform={
                "operation": "parse_only",
                "version": "1.0.0",
                "encoding": "xml_declared",
                "newline": "preserve",
                "namespace_mode": "cda_local_name",
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one CDA document path against caller-held XML."""

        payload = _coerce_bytes(source)
        if payload is None or locator.location_type != "document_path":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, payload):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        if _UNSAFE_XML_RE.search(payload):
            return StoreResult.outcome(StoreState.DENIED, "cda_declaration_unsafe")
        try:
            root = ET.fromstring(payload)
            return StoreResult.success(
                _resolve_cda_path(root, str(locator.location["path"]))
            )
        except (ET.ParseError, KeyError, IndexError, TypeError, ValueError):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_unresolvable")


class DelimitedTableEvidenceAdapter:
    """Adapt CSV or another single-character delimited UTF-8 table."""

    parser_id = "openmed.evidence.csv"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        delimiter: str = ",",
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        if (
            not isinstance(delimiter, str)
            or len(delimiter) != 1
            or delimiter in {'"', "\r", "\n"}
        ):
            raise ValueError("delimiter must be one safe character")
        self.delimiter = delimiter
        self.source_version = "rfc4180" if delimiter == "," else "delimited1"
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Parse every logical cell with one-based row/column coordinates."""

        coerced = _coerce_utf8(source, strip_bom=True)
        if coerced is None:
            return _quarantine(
                _coerce_bytes(source) or b"",
                context,
                source_format="csv",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="csv_source_invalid",
            )
        payload, text = coerced
        limited = _limit_quarantine(
            payload,
            context,
            source_format="csv",
            source_version=self.source_version,
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        try:
            rows = _csv_rows(text, self.delimiter)
        except csv.Error:
            return _quarantine(
                payload,
                context,
                source_format="csv",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="csv_malformed",
            )
        field_count = sum(len(row) for row in rows)
        if field_count > self.max_fields:
            return _field_limit_quarantine(
                payload,
                context,
                source_format="csv",
                source_version=self.source_version,
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=field_count,
            )
        specs: tuple[LocatorSpec, ...] = tuple(
            ("table_cell", {"row": row_index, "column": column_index}, {})
            for row_index, row in enumerate(rows, start=1)
            for column_index, _value in enumerate(row, start=1)
        )
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="csv",
            source_version=self.source_version,
            media_type="text/csv",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="table_cell_v1",
            normalization_transform={
                "operation": "logical_cell_parse",
                "version": "1.0.0",
                "encoding": "utf8_sig",
                "newline": "preserve",
                "delimiter": _delimiter_name(self.delimiter),
                "index_base": 1,
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one logical CSV cell, including quoted newlines."""

        coerced = _coerce_utf8(source, strip_bom=True)
        if coerced is None or locator.location_type != "table_cell":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, coerced[0]):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        try:
            rows = _csv_rows(coerced[1], self.delimiter)
            row = int(locator.location["row"])
            column = int(locator.location["column"])
            return StoreResult.success(rows[row - 1][column - 1])
        except (csv.Error, IndexError, TypeError, ValueError):
            return StoreResult.outcome(StoreState.CONFLICT, "locator_unresolvable")


class XLSXEvidenceAdapter:
    """Adapt XLSX cells with opaque sheet-order coordinates."""

    parser_id = "openmed.evidence.xlsx"
    parser_version = "1.0.0"

    def __init__(
        self,
        *,
        max_source_bytes: int = DEFAULT_MAX_SOURCE_BYTES,
        max_fields: int = DEFAULT_MAX_FIELDS,
    ) -> None:
        self.max_source_bytes = _positive_limit(max_source_bytes, "max_source_bytes")
        self.max_fields = _positive_limit(max_fields, "max_fields")

    def adapt(
        self,
        source: Any,
        context: EvidenceAdapterContext,
    ) -> StoreResult[EvidenceAdapterOutcome]:
        """Parse non-empty workbook cells without exposing sheet titles."""

        payload = _coerce_bytes(source)
        if payload is None:
            payload = b""
            return _quarantine(
                payload,
                context,
                source_format="xlsx",
                source_version="ecma376",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="xlsx_source_invalid",
            )
        limited = _limit_quarantine(
            payload,
            context,
            source_format="xlsx",
            source_version="ecma376",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            max_source_bytes=self.max_source_bytes,
        )
        if limited is not None:
            return limited
        if importlib.util.find_spec("openpyxl") is None:
            return _quarantine(
                payload,
                context,
                source_format="xlsx",
                source_version="ecma376",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.UNSUPPORTED,
                classification="unsupported",
                reason_code="xlsx_dependency_missing",
            )
        workbook = None
        try:
            workbook = _load_workbook(payload)
            cells = _xlsx_cells(workbook, self.max_fields)
        except OverflowError:
            return _field_limit_quarantine(
                payload,
                context,
                source_format="xlsx",
                source_version="ecma376",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                candidate_count=self.max_fields + 1,
            )
        except Exception:
            return _quarantine(
                payload,
                context,
                source_format="xlsx",
                source_version="ecma376",
                parser_id=self.parser_id,
                parser_version=self.parser_version,
                state=StoreState.FAILURE,
                classification="malformed",
                reason_code="xlsx_malformed",
            )
        finally:
            if workbook is not None:
                workbook.close()
        specs: tuple[LocatorSpec, ...] = tuple(
            (
                "table_cell",
                {"sheet": f"sheet_{sheet_index:04d}", "row": row, "column": column},
                {},
            )
            for sheet_index, row, column in cells
        )
        return _success(
            source_bytes=payload,
            artifact_bytes=payload,
            context=context,
            source_format="xlsx",
            source_version="ecma376",
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            parser_id=self.parser_id,
            parser_version=self.parser_version,
            coordinate_convention="table_cell_v1",
            normalization_transform={
                "operation": "logical_cell_parse",
                "version": "1.0.0",
                "sheet_mode": "opaque_order",
                "formula_mode": "formula",
                "index_base": 1,
            },
            locator_specs=specs,
        )

    def resolve(
        self,
        source: Any,
        locator: EvidenceLocator,
    ) -> StoreResult[Any]:
        """Resolve one workbook cell by opaque sheet order and A1 coordinates."""

        payload = _coerce_bytes(source)
        if payload is None or locator.location_type != "table_cell":
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        if not _locator_matches_source(locator, payload):
            return StoreResult.outcome(StoreState.CONFLICT, "source_digest_mismatch")
        sheet = locator.location.get("sheet")
        match = re.fullmatch(r"sheet_([0-9]{4})", str(sheet))
        if match is None or importlib.util.find_spec("openpyxl") is None:
            return StoreResult.outcome(StoreState.UNSUPPORTED, "locator_unsupported")
        workbook = None
        try:
            workbook = _load_workbook(payload)
            sheet_index = int(match.group(1))
            worksheet = workbook.worksheets[sheet_index - 1]
            row = int(locator.location["row"])
            column = int(locator.location["column"])
            return StoreResult.success(worksheet.cell(row=row, column=column).value)
        except Exception:
            return StoreResult.outcome(StoreState.CONFLICT, "locator_unresolvable")
        finally:
            if workbook is not None:
                workbook.close()


def _success(
    *,
    source_bytes: bytes,
    artifact_bytes: bytes,
    context: EvidenceAdapterContext,
    source_format: str,
    source_version: str,
    media_type: str,
    parser_id: str,
    parser_version: str,
    coordinate_convention: str,
    normalization_transform: Mapping[str, Any],
    locator_specs: tuple[LocatorSpec, ...],
) -> StoreResult[EvidenceAdapterOutcome]:
    source_digest = sha256_digest(source_bytes)
    artifact_digest = sha256_digest(artifact_bytes)
    artifact = ClinicalArtifact(
        artifact_id=derived_opaque_id(
            "artifact",
            source_digest,
            source_format,
            source_version,
            parser_id,
            parser_version,
            context.source_id,
            context.subject_id,
            context.encounter_id,
            context.recorded_at,
        ),
        artifact_type=source_format,
        media_type=media_type,
        content_hash=artifact_digest,
        byte_size=len(artifact_bytes),
        source_id=context.source_id,
        recorded_at=context.recorded_at,
        subject_id=context.subject_id,
        encounter_id=context.encounter_id,
        attributes={
            "coordinate_convention": coordinate_convention,
            "normalization_transform": dict(normalization_transform),
            "parser_id": parser_id,
            "parser_version": parser_version,
            "source_digest": source_digest,
            "source_format": source_format,
            "source_version": source_version,
        },
    )
    locators = tuple(
        EvidenceLocator(
            locator_id=derived_opaque_id(
                "evidence",
                artifact.artifact_id,
                location_type,
                dict(location),
                {"source_digest": source_digest, **dict(transform)},
            ),
            artifact_id=artifact.artifact_id,
            location_type=location_type,
            location=location,
            transform={"source_digest": source_digest, **dict(transform)},
            schema_version=("1.1.0" if location_type == "document_path" else "1.0.0"),
        )
        for location_type, location, transform in locator_specs
    )
    result = StructuredEvidenceResult(
        result_id=derived_opaque_id(
            "evidence_result",
            artifact.artifact_id,
            tuple(sorted(item.locator_id for item in locators)),
            parser_version,
        ),
        artifact=artifact,
        locators=locators,
        source_digest=source_digest,
        source_byte_size=len(source_bytes),
        source_format=source_format,
        source_version=source_version,
        parser_id=parser_id,
        parser_version=parser_version,
        coordinate_convention=coordinate_convention,
        normalization_transform=normalization_transform,
    )
    return StoreResult.success(result)


def _quarantine(
    source_bytes: bytes,
    context: EvidenceAdapterContext,
    *,
    source_format: str,
    source_version: str,
    parser_id: str,
    parser_version: str,
    state: StoreState,
    classification: str,
    reason_code: str,
    candidate_count: int = 0,
) -> StoreResult[EvidenceAdapterOutcome]:
    digest = sha256_digest(source_bytes)
    quarantine = StructuredEvidenceQuarantine(
        quarantine_id=derived_opaque_id(
            "evidence_quarantine",
            digest,
            source_format,
            source_version,
            parser_id,
            parser_version,
            classification,
            reason_code,
        ),
        source_digest=digest,
        source_byte_size=len(source_bytes),
        source_format=source_format,
        source_version=_safe_version(source_version),
        parser_id=parser_id,
        parser_version=parser_version,
        classification=classification,
        reason_code=reason_code,
        candidate_count=candidate_count,
        failure_count=1,
        created_at=context.recorded_at,
    )
    return StoreResult.outcome(state, reason_code, value=quarantine)


def _limit_quarantine(
    source_bytes: bytes,
    context: EvidenceAdapterContext,
    *,
    source_format: str,
    source_version: str,
    parser_id: str,
    parser_version: str,
    max_source_bytes: int,
) -> StoreResult[EvidenceAdapterOutcome] | None:
    if len(source_bytes) <= max_source_bytes:
        return None
    return _quarantine(
        source_bytes,
        context,
        source_format=source_format,
        source_version=source_version,
        parser_id=parser_id,
        parser_version=parser_version,
        state=StoreState.DENIED,
        classification="unsafe",
        reason_code="source_limit_exceeded",
    )


def _field_limit_quarantine(
    source_bytes: bytes,
    context: EvidenceAdapterContext,
    *,
    source_format: str,
    source_version: str,
    parser_id: str,
    parser_version: str,
    candidate_count: int,
) -> StoreResult[EvidenceAdapterOutcome]:
    return _quarantine(
        source_bytes,
        context,
        source_format=source_format,
        source_version=source_version,
        parser_id=parser_id,
        parser_version=parser_version,
        state=StoreState.PARTIAL,
        classification="partial",
        reason_code="field_limit_exceeded",
        candidate_count=candidate_count,
    )


def _coerce_bytes(source: Any) -> bytes | None:
    if isinstance(source, str):
        try:
            return source.encode("utf-8")
        except UnicodeEncodeError:
            return None
    if isinstance(source, bytes):
        return source
    if isinstance(source, (bytearray, memoryview)):
        return bytes(source)
    return None


def _coerce_utf8(
    source: Any,
    *,
    strip_bom: bool = False,
) -> tuple[bytes, str] | None:
    payload = _coerce_bytes(source)
    if payload is None:
        return None
    try:
        text = payload.decode("utf-8-sig" if strip_bom else "utf-8")
    except UnicodeDecodeError:
        return None
    return payload, text


def _strict_json(text: str) -> Any:
    return json.loads(
        text,
        object_pairs_hook=_reject_duplicate_keys,
        parse_constant=_reject_json_constant,
    )


def _reject_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("duplicate key")
        result[key] = value
    return result


def _reject_json_constant(value: str) -> None:
    raise ValueError("non-finite number")


def _json_leaf_pointers(value: Any, pointer: str = "") -> tuple[str, ...]:
    pointers: list[str] = []

    def visit(item: Any, path: str) -> None:
        if isinstance(item, Mapping):
            for key in sorted(item):
                if not isinstance(key, str) or _FHIR_KEY_RE.fullmatch(key) is None:
                    raise ValueError("FHIR property key is unsupported")
                escaped = key.replace("~", "~0").replace("/", "~1")
                visit(item[key], f"{path}/{escaped}")
            return
        if isinstance(item, list):
            for index, child in enumerate(item):
                visit(child, f"{path}/{index}")
            return
        pointers.append(path)

    visit(value, pointer)
    return tuple(pointers)


def _resolve_json_pointer(value: Any, pointer: str) -> Any:
    if pointer == "":
        return value
    current = value
    for raw_token in pointer.split("/")[1:]:
        token = raw_token.replace("~1", "/").replace("~0", "~")
        if isinstance(current, Mapping):
            current = current[token]
        elif isinstance(current, list):
            current = current[int(token)]
        else:
            raise TypeError("pointer crosses a scalar")
    return current


def _hl7_component_paths(message: HL7Message) -> tuple[str, ...]:
    occurrences: dict[str, int] = {}
    paths: list[str] = []
    for segment in message.segments:
        occurrence = occurrences.get(segment.name, 0) + 1
        occurrences[segment.name] = occurrence
        if segment.name == "MSH":
            paths.extend((f"MSH.1.1[{occurrence}]", f"MSH.2.1[{occurrence}]"))
            fields = enumerate(segment.fields[1:], start=3)
        else:
            fields = enumerate(segment.fields, start=1)
        for field_position, field_value in fields:
            if not field_value:
                continue
            if (
                segment.encoding.repetition in field_value
                or segment.encoding.subcomponent in field_value
            ):
                raise ValueError("field requires unsupported coordinate dimension")
            components = field_value.split(segment.encoding.component)
            for component_position, component in enumerate(components, start=1):
                if component:
                    paths.append(
                        f"{segment.name}.{field_position}.{component_position}"
                        f"[{occurrence}]"
                    )
    return tuple(paths)


def _resolve_hl7_component(message: HL7Message, path: str) -> str:
    match = _HL7_PATH_RE.fullmatch(path)
    if match is None:
        raise ValueError("invalid path")
    segment_name = match.group("segment")
    occurrence = int(match.group("occurrence"))
    field_position = int(match.group("field"))
    component_position = int(match.group("component"))
    matching = [item for item in message.segments if item.name == segment_name]
    segment = matching[occurrence - 1]
    if segment_name == "MSH" and field_position == 1:
        value = segment.encoding.field
    elif segment_name == "MSH" and field_position == 2:
        value = segment.encoding.as_msh_2()
    else:
        value = segment.get_field(field_position)
        if value is None:
            raise IndexError("field absent")
    if field_position in {1, 2} and segment_name == "MSH":
        components = [value]
    else:
        if (
            segment.encoding.repetition in value
            or segment.encoding.subcomponent in value
        ):
            raise ValueError("ambiguous field")
        components = value.split(segment.encoding.component)
    return components[component_position - 1]


def _cda_value_paths(root: ET.Element) -> tuple[tuple[str, int | None], ...]:
    result: list[tuple[str, int | None]] = []
    section_counter = 0

    def visit(element: ET.Element, path: str, section: int | None) -> None:
        nonlocal section_counter
        namespace = _namespace(element.tag)
        if namespace not in {"", _CDA_NAMESPACE}:
            raise ValueError("element namespace is unsupported")
        if _local_name(element.tag) == "section":
            section_counter += 1
            section = section_counter
        attribute_names: dict[str, int] = {}
        for attribute in element.attrib:
            namespace = _namespace(attribute)
            if namespace not in {"", _CDA_NAMESPACE, _XSI_NAMESPACE}:
                raise ValueError("attribute namespace is unsupported")
            local = _local_name(attribute)
            attribute_names[local] = attribute_names.get(local, 0) + 1
        if any(count > 1 for count in attribute_names.values()):
            raise ValueError("attribute coordinate is ambiguous")
        for local in sorted(attribute_names):
            result.append((f"{path}/@{local}", section))
        if (
            element.text is not None
            and not element.text.isspace()
            and element.text != ""
        ):
            result.append((f"{path}/text()", section))
        occurrences: dict[str, int] = {}
        for child in list(element):
            local = _local_name(child.tag)
            occurrences[local] = occurrences.get(local, 0) + 1
            child_path = f"{path}/{local}[{occurrences[local]}]"
            visit(child, child_path, section)
            if child.tail is not None and not child.tail.isspace() and child.tail != "":
                result.append((f"{child_path}/tail()", section))

    root_path = f"/{_local_name(root.tag)}[1]"
    visit(root, root_path, None)
    return tuple(result)


def _resolve_cda_path(root: ET.Element, path: str) -> str:
    pieces = path.split("/")[1:]
    if not pieces:
        raise ValueError("empty path")
    terminal = None
    if pieces[-1] in {"text()", "tail()"} or pieces[-1].startswith("@"):
        terminal = pieces.pop()
    root_match = _XML_SEGMENT_RE.fullmatch(pieces.pop(0))
    if (
        root_match is None
        or root_match.group("name") != _local_name(root.tag)
        or int(root_match.group("index")) != 1
    ):
        raise ValueError("root differs")
    current = root
    for piece in pieces:
        match = _XML_SEGMENT_RE.fullmatch(piece)
        if match is None:
            raise ValueError("invalid element path")
        matching = [
            child
            for child in list(current)
            if _local_name(child.tag) == match.group("name")
        ]
        current = matching[int(match.group("index")) - 1]
    if terminal == "text()":
        return current.text or ""
    if terminal == "tail()":
        return current.tail or ""
    if terminal is not None and terminal.startswith("@"):
        name = terminal[1:]
        matches = [
            value for key, value in current.attrib.items() if _local_name(key) == name
        ]
        if len(matches) != 1:
            raise KeyError("attribute differs")
        return matches[0]
    return current.text or ""


def _local_name(name: str) -> str:
    return name.rsplit("}", 1)[-1] if name.startswith("{") else name


def _namespace(name: str) -> str:
    return name[1:].split("}", 1)[0] if name.startswith("{") else ""


def _csv_rows(text: str, delimiter: str) -> list[list[str]]:
    return list(
        csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    )


def _delimiter_name(delimiter: str) -> str:
    return {
        ",": "comma",
        "\t": "tab",
        ";": "semicolon",
        "|": "pipe",
    }.get(delimiter, "custom")


def _load_workbook(payload: bytes) -> Any:
    from openpyxl import load_workbook

    return load_workbook(
        io.BytesIO(payload),
        read_only=True,
        data_only=False,
        keep_links=False,
    )


def _xlsx_cells(workbook: Any, max_fields: int) -> tuple[tuple[int, int, int], ...]:
    result: list[tuple[int, int, int]] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                result.append((sheet_index, cell.row, cell.column))
                if len(result) > max_fields:
                    raise OverflowError("field limit")
    return tuple(result)


def _positive_limit(value: Any, field_name: str) -> int:
    if type(value) is not int or value < 1:
        raise ValueError(f"{field_name} must be positive")
    return value


def _safe_version(value: Any) -> str:
    if isinstance(value, str) and _SAFE_VERSION_RE.fullmatch(value) is not None:
        return value
    return "unknown"


def _locator_matches_source(locator: EvidenceLocator, source_bytes: bytes) -> bool:
    return locator.transform.get("source_digest") == sha256_digest(source_bytes)


__all__ = [
    "DEFAULT_MAX_FIELDS",
    "DEFAULT_MAX_SOURCE_BYTES",
    "CDAEvidenceAdapter",
    "DelimitedTableEvidenceAdapter",
    "EvidenceAdapterContext",
    "EvidenceAdapterOutcome",
    "ExistingDocumentEvidenceAdapter",
    "ExistingDocumentInput",
    "FHIRR4EvidenceAdapter",
    "HL7V2EvidenceAdapter",
    "StructuredEvidenceAdapter",
    "TextEvidenceAdapter",
    "XLSXEvidenceAdapter",
]
